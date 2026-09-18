"""DOCX 公文质量检查器。"""

from pathlib import Path
from zipfile import ZipFile, BadZipFile
from docx import Document
from docx.oxml.ns import qn

def inspect(path: str):
    errors, warnings, info = [], [], []
    p = Path(path)
    if not p.exists():
        return [f"文件不存在：{path}"], [], []
    try:
        with ZipFile(p) as z:
            names = set(z.namelist())
            for must in ("[Content_Types].xml", "word/document.xml"):
                if must not in names:
                    errors.append(f"DOCX结构缺失：{must}")
    except BadZipFile:
        errors.append("文件不是有效的 DOCX 压缩结构")
        return errors, warnings, info

    try:
        doc = Document(str(p))
    except Exception as e:
        errors.append(f"Word 文档无法读取：{e}")
        return errors, warnings, info

    info.append(f"节数量：{len(doc.sections)}")
    for i, sec in enumerate(doc.sections):
        mm = lambda x: round(x / 36000, 1) if x is not None else None
        vals = {
            "宽": mm(sec.page_width), "高": mm(sec.page_height),
            "上": mm(sec.top_margin), "下": mm(sec.bottom_margin),
            "左": mm(sec.left_margin), "右": mm(sec.right_margin),
        }
        info.append(f"第{i+1}节：{vals}")
        if (abs(vals["宽"] - 210) > 1 or abs(vals["高"] - 297) > 1):
            warnings.append(f"第{i+1}节纸张不是A4：{vals}")

    settings = doc.settings.element
    uf = settings.find(qn("w:updateFields"))
    if uf is None or uf.get(qn("w:val")) != "true":
        warnings.append("未启用 Word 打开时字段自动更新")

    if not any(p.text.strip() for p in doc.paragraphs):
        errors.append("文档没有可见正文")

    blank = 0
    long_blank = 0
    for p in doc.paragraphs:
        if not p.text.strip():
            blank += 1
            if blank > 5:
                long_blank += 1
        else:
            blank = 0
    if long_blank:
        warnings.append("检测到连续空段，建议检查是否使用空行制造分页")

    return errors, warnings, info


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    args = ap.parse_args()
    errors, warnings, info = inspect(args.file)
    print("【硬错误】")
    print("\n".join(errors) if errors else "无")
    print("\n【风险项】")
    print("\n".join(warnings) if warnings else "无")
    print("\n【检查信息】")
    print("\n".join(info))
    raise SystemExit(1 if errors else 0)
